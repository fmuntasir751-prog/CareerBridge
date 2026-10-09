import hashlib
import json
import os
import re
from pathlib import Path

from django.test import client
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader


MAX_RESUME_SIZE = 5 * 1024 * 1024
MAX_RESUME_TEXT_LENGTH = 20000

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
}

SKILL_KEYWORDS = {
    "React": ["react", "react.js", "reactjs"],
    "JavaScript": [
        "javascript",
        "js",
        "es6",
    ],
    "TypeScript": [
        "typescript",
        "ts",
    ],
    "HTML": ["html", "html5"],
    "CSS": ["css", "css3"],
    "Python": ["python"],
    "Django": [
        "django",
        "django rest framework",
        "drf",
    ],
    "REST API": [
        "rest api",
        "restful api",
    ],
    "PostgreSQL": [
        "postgresql",
        "postgres",
    ],
    "MySQL": ["mysql"],
    "Git": ["git", "github"],
    "Docker": ["docker"],
    "Node.js": [
        "node.js",
        "nodejs",
    ],
    "Vue": [
        "vue",
        "vue.js",
        "vuejs",
    ],
    "Angular": ["angular"],
    "Next.js": [
        "next.js",
        "nextjs",
    ],
    "Tailwind CSS": [
        "tailwind",
        "tailwind css",
    ],
    "Bootstrap": ["bootstrap"],
    "AWS": [
        "aws",
        "amazon web services",
    ],
    "Figma": ["figma"],
    "Java": ["java"],
    "C++": ["c++"],
}


class ResumeAnalysisError(Exception):
    pass


def calculate_resume_hash(resume):
    digest = hashlib.sha256()

    resume.open("rb")

    try:
        while True:
            chunk = resume.read(8192)

            if not chunk:
                break

            digest.update(chunk)
    finally:
        resume.seek(0)

    return digest.hexdigest()


def validate_resume(resume):
    extension = Path(
        resume.name,
    ).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeAnalysisError(
            "Only PDF and DOCX resumes are supported."
        )

    try:
        file_size = resume.size
    except (AttributeError, OSError):
        file_size = None

    if (
        file_size is not None
        and file_size > MAX_RESUME_SIZE
    ):
        raise ResumeAnalysisError(
            "The resume must be 5 MB or smaller."
        )

    return extension


def extract_pdf_text(resume):
    resume.open("rb")

    try:
        reader = PdfReader(resume)

        text_parts = []

        for page in reader.pages:
            page_text = page.extract_text() or ""

            if page_text.strip():
                text_parts.append(
                    page_text.strip(),
                )

        return "\n".join(text_parts)
    except Exception as error:
        raise ResumeAnalysisError(
            "The PDF resume could not be read."
        ) from error
    finally:
        resume.seek(0)


def extract_docx_text(resume):
    resume.open("rb")

    try:
        document = Document(resume)

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if values:
                    paragraphs.append(
                        " | ".join(values),
                    )

        return "\n".join(paragraphs)
    except Exception as error:
        raise ResumeAnalysisError(
            "The DOCX resume could not be read."
        ) from error
    finally:
        resume.seek(0)


def extract_resume_text(resume):
    extension = validate_resume(resume)

    if extension == ".pdf":
        text = extract_pdf_text(resume)
    else:
        text = extract_docx_text(resume)

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    ).strip()

    if len(text) < 80:
        raise ResumeAnalysisError(
            "Not enough readable text was found "
            "in the resume. Scanned image PDFs "
            "are not currently supported."
        )

    return text[:MAX_RESUME_TEXT_LENGTH]


def contains_keyword(text, keyword):
    pattern = (
        r"(?<![a-zA-Z0-9])"
        + re.escape(keyword.lower())
        + r"(?![a-zA-Z0-9])"
    )

    return (
        re.search(pattern, text.lower())
        is not None
    )


def detect_skills(text):
    detected = []

    for skill, aliases in SKILL_KEYWORDS.items():
        if any(
            contains_keyword(text, alias)
            for alias in aliases
        ):
            detected.append(skill)

    return sorted(detected)


def create_local_analysis(
    resume_text,
    language,
):
    text = resume_text.lower()
    detected_skills = detect_skills(
        resume_text,
    )

    section_checks = {
        "contact": any(
            value in text
            for value in (
                "@",
                "email",
                "phone",
                "電話",
                "メール",
            )
        ),
        "summary": any(
            value in text
            for value in (
                "summary",
                "profile",
                "objective",
                "自己紹介",
                "概要",
                "志望動機",
            )
        ),
        "skills": bool(
            detected_skills
        )
        or any(
            value in text
            for value in (
                "skills",
                "technical skills",
                "スキル",
                "技術",
            )
        ),
        "experience": any(
            value in text
            for value in (
                "experience",
                "employment",
                "work history",
                "職歴",
                "経験",
            )
        ),
        "education": any(
            value in text
            for value in (
                "education",
                "university",
                "school",
                "学歴",
                "大学",
                "専門学校",
            )
        ),
        "projects": any(
            value in text
            for value in (
                "project",
                "portfolio",
                "github",
                "プロジェクト",
                "ポートフォリオ",
            )
        ),
        "languages": any(
            value in text
            for value in (
                "language",
                "jlpt",
                "toeic",
                "japanese",
                "english",
                "言語",
                "日本語",
                "英語",
            )
        ),
    }

    section_weights = {
        "contact": 10,
        "summary": 10,
        "skills": 20,
        "experience": 20,
        "education": 15,
        "projects": 15,
        "languages": 10,
    }

    section_scores = {
        section: (
            weight
            if section_checks[section]
            else 0
        )
        for section, weight
        in section_weights.items()
    }

    overall_score = sum(
        section_scores.values()
    )

    missing_sections = [
        section
        for section, found
        in section_checks.items()
        if not found
    ]

    if language == "ja":
        strengths = []

        if detected_skills:
            strengths.append(
                "技術スキルが明確に記載されています。"
            )

        if section_checks["projects"]:
            strengths.append(
                "プロジェクト経験が記載されています。"
            )

        if section_checks["languages"]:
            strengths.append(
                "語学力や資格が確認できます。"
            )

        if not strengths:
            strengths.append(
                "履歴書から基本情報を確認できました。"
            )

        recommendations = []

        recommendation_map = {
            "contact": (
                "メールアドレスと連絡先を"
                "分かりやすく記載してください。"
            ),
            "summary": (
                "冒頭に短い自己紹介と"
                "希望職種を追加してください。"
            ),
            "skills": (
                "技術スキルをカテゴリー別に"
                "記載してください。"
            ),
            "experience": (
                "職歴や実務経験を成果とともに"
                "記載してください。"
            ),
            "education": (
                "学校名、専攻、卒業予定年を"
                "追加してください。"
            ),
            "projects": (
                "プロジェクト、使用技術、"
                "GitHub URLを追加してください。"
            ),
            "languages": (
                "JLPT、TOEICなどの"
                "語学資格を追加してください。"
            ),
        }

        for section in missing_sections:
            recommendations.append(
                recommendation_map[section]
            )

        summary = (
            "履歴書の構成、技術スキル、"
            "経験、学歴を確認しました。"
        )
    else:
        strengths = []

        if detected_skills:
            strengths.append(
                "Technical skills are clearly included."
            )

        if section_checks["projects"]:
            strengths.append(
                "Project experience is included."
            )

        if section_checks["languages"]:
            strengths.append(
                "Language skills or certifications are included."
            )

        if not strengths:
            strengths.append(
                "Basic resume information was detected."
            )

        recommendations = []

        recommendation_map = {
            "contact": (
                "Add a clear email address "
                "and contact information."
            ),
            "summary": (
                "Add a short professional summary "
                "and desired position."
            ),
            "skills": (
                "Organize technical skills "
                "into clear categories."
            ),
            "experience": (
                "Describe work experience "
                "with measurable achievements."
            ),
            "education": (
                "Add your school, major "
                "and expected graduation year."
            ),
            "projects": (
                "Add projects, technologies "
                "and GitHub links."
            ),
            "languages": (
                "Add language qualifications "
                "such as JLPT or TOEIC."
            ),
        }

        for section in missing_sections:
            recommendations.append(
                recommendation_map[section]
            )

        summary = (
            "The resume was reviewed for structure, "
            "technical skills, experience and education."
        )

    return {
        "overall_score": overall_score,
        "summary": summary,
        "detected_skills": detected_skills,
        "strengths": strengths,
        "missing_sections": missing_sections,
        "recommendations": recommendations,
        "section_scores": section_scores,
        "ai_generated": False,
    }


def clean_json_response(value):
    value = value.strip()

    value = re.sub(
        r"^```(?:json)?\s*",
        "",
        value,
        flags=re.IGNORECASE,
    )

    value = re.sub(
        r"\s*```$",
        "",
        value,
    )

    return value.strip()


def normalize_analysis(data, fallback):
    overall_score = data.get(
        "overall_score",
        fallback["overall_score"],
    )

    try:
        overall_score = int(
            overall_score,
        )
    except (TypeError, ValueError):
        overall_score = (
            fallback["overall_score"]
        )

    overall_score = max(
        0,
        min(100, overall_score),
    )

    result = {
        "overall_score": overall_score,
        "summary": str(
            data.get(
                "summary",
                fallback["summary"],
            )
        )[:1000],
        "detected_skills": list(
            data.get(
                "detected_skills",
                fallback["detected_skills"],
            )
        )[:30],
        "strengths": list(
            data.get(
                "strengths",
                fallback["strengths"],
            )
        )[:10],
        "missing_sections": list(
            data.get(
                "missing_sections",
                fallback["missing_sections"],
            )
        )[:10],
        "recommendations": list(
            data.get(
                "recommendations",
                fallback["recommendations"],
            )
        )[:10],
        "section_scores": fallback[
            "section_scores"
        ],
        "ai_generated": True,
    }

    provided_scores = data.get(
        "section_scores",
    )

    if isinstance(
        provided_scores,
        dict,
    ):
        normalized_scores = {}

        for name, score in provided_scores.items():
            try:
                normalized_scores[str(name)] = max(
                    0,
                    min(100, int(score)),
                )
            except (TypeError, ValueError):
                continue

        if normalized_scores:
            result["section_scores"] = (
                normalized_scores
            )

    return result


def generate_ai_analysis(
    resume_text,
    language,
    profile,
):
    fallback = create_local_analysis(
        resume_text,
        language,
    )

    api_key = os.getenv(
        "GEMINI_API_KEY",
    )

    if not api_key:
        return fallback

    response_language = (
        "Japanese"
        if language == "ja"
        else "English"
    )

    prompt = f"""
You are a professional resume reviewer for
CareerBridge, a job platform for international
students seeking employment in Japan.

Analyze the resume in {response_language}.

Return valid JSON only with this exact structure:
{{
  "overall_score": 0,
  "summary": "",
  "detected_skills": [],
  "strengths": [],
  "missing_sections": [],
  "recommendations": [],
  "section_scores": {{
    "contact": 0,
    "summary": 0,
    "skills": 0,
    "experience": 0,
    "education": 0,
    "projects": 0,
    "languages": 0
  }}
}}

Rules:
- overall_score must be between 0 and 100.
- Give practical, beginner-friendly feedback.
- Do not invent experience, skills or qualifications.
- Do not include sensitive personal data in the output.
- Each list must contain short strings.
- section_scores must use values from 0 to 100.
- Focus on applications to companies in Japan.
- Consider the desired job title:
  {profile.desired_job_title or "Not specified"}

Resume text:
{resume_text}
"""

    try:
        client = genai.Client(
    api_key=api_key,
    http_options=types.HttpOptions(
        timeout=10000,
        retry_options=types.HttpRetryOptions(
            attempts=1,
        ),
    ),
)

        interaction = (
            client.interactions.create(
                model="gemini-3.6-flash",
                input=prompt,
            )
        )

        response_text = (
            interaction.output_text or ""
        )

        if not response_text:
            return fallback

        data = json.loads(
            clean_json_response(
                response_text,
            )
        )

        if not isinstance(data, dict):
            return fallback

        return normalize_analysis(
            data,
            fallback,
        )
    except Exception as error:
        print(
            f"Resume analysis failed: {error}",
            flush=True,
        )

        return fallback